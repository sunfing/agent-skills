#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path

FOREST_MAP = {"011":"重点公益林","012":"一般公益林","021":"重点商品林","022":"一般商品林"}
GRADE_MAP = {"1":"Ⅰ","2":"Ⅱ","3":"Ⅲ","4":"Ⅳ","I":"Ⅰ","II":"Ⅱ","III":"Ⅲ","IV":"Ⅳ"}
ALIASES = {
    "parcel":["序号","地块号","小班号","DKBH","序号1"],
    "area":["面积","小班面积","TBMJ","XBMJ","Shape_Area"],
    "land":["地类","DI_LEI","DLMC"],
    "forest":["森林类","森林类别","SEN_LIN_LB","SBN_LIN_LB"],
    "grade":["保护等","保护等级","BH_DJ"],
    "point_x":["POINT_X"], "point_y":["POINT_Y"],
}

class ConversionError(RuntimeError):
    pass

@dataclass
class PointRow:
    source_row:int; parcel:object; area:float|None; land:str; forest:str; grade:str
    point_x:float; point_y:float

@dataclass
class OutputRow:
    source_row:int; parcel:object; ring:int; corner:str; area:float|None
    land:str; forest:str; grade:str; point_x:float; point_y:float
    @property
    def northing_x(self): return round(self.point_y, 3)
    @property
    def easting_y(self): return round(self.point_x, 3)

def clean(v):
    if v is None: return ""
    if isinstance(v,float) and math.isnan(v): return ""
    s=str(v).strip()
    return "" if s.lower() in {"nan","none","null","<空>"} else s

def norm(v):
    return re.sub(r"\s+","",clean(v).replace("\ufeff","")).upper()

def as_float(v,name,row):
    try:
        if clean(v)=="": raise ValueError
        return float(v)
    except Exception as e:
        raise ConversionError(f"第 {row} 行字段 {name} 不是有效数值：{v!r}") from e

def opt_float(v):
    try: return None if clean(v)=="" else float(v)
    except Exception: return None

def parcel_value(v):
    s=clean(v)
    if not s: return ""
    try:
        f=float(s)
        return int(f) if f.is_integer() else s
    except Exception: return s

def std_forest(v):
    s=clean(v); key=s.zfill(3) if s.isdigit() else s
    return FOREST_MAP.get(key,s)

def std_grade(v):
    s=clean(v)
    if s.endswith(".0"): s=s[:-2]
    return GRADE_MAP.get(s.upper(),s)

def find_header(headers,candidates):
    hs=[norm(h) for h in headers]
    for c in candidates:
        c=norm(c)
        if c in hs: return hs.index(c)
    return None

def read_table(path):
    ext=path.suffix.lower()
    if ext==".csv":
        last=None
        for enc in ("utf-8-sig","gb18030","utf-8"):
            try:
                with path.open("r",encoding=enc,newline="") as f: rows=list(csv.reader(f))
                if not rows: raise ConversionError("CSV 文件为空。")
                return rows[0],rows[1:]
            except UnicodeDecodeError as e: last=e
        raise ConversionError(f"无法识别 CSV 编码：{last}")
    if ext==".xls":
        try: import xlrd
        except ImportError as e: raise ConversionError("读取 .xls 需要 xlrd，请安装 requirements.txt。") from e
        book=xlrd.open_workbook(str(path),on_demand=True)
        sh=book.sheet_by_index(0)
        if sh.nrows<2: raise ConversionError("XLS 文件没有数据行。")
        return [str(x) for x in sh.row_values(0)],[sh.row_values(i) for i in range(1,sh.nrows)]
    if ext==".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as e: raise ConversionError("读取 .xlsx 需要 openpyxl，请安装 requirements.txt。") from e
        wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
        it=ws.iter_rows(values_only=True)
        try: headers=list(next(it))
        except StopIteration: raise ConversionError("XLSX 文件为空。")
        return ["" if x is None else str(x) for x in headers],[list(r) for r in it]
    raise ConversionError("仅支持 .xls、.xlsx、.csv 初表。")

def parse_rows(headers,raw,parcel_field=None):
    fields={}
    fields["parcel"]=find_header(headers,[parcel_field] if parcel_field else ALIASES["parcel"])
    for k in ("area","land","forest","grade","point_x","point_y"):
        fields[k]=find_header(headers,ALIASES[k])
    missing=[n for k,n in (("parcel","地块字段"),("point_x","POINT_X"),("point_y","POINT_Y")) if fields[k] is None]
    if missing: raise ConversionError("缺少必要字段："+"、".join(missing))
    resolved={k:(str(headers[i]) if i is not None else "") for k,i in fields.items()}
    def get(r,i): return None if i is None or i>=len(r) else r[i]
    out=[]
    for pos,r in enumerate(raw,start=2):
        pv,px,py=get(r,fields["parcel"]),get(r,fields["point_x"]),get(r,fields["point_y"])
        if clean(pv)==clean(px)==clean(py)=="": continue
        p=parcel_value(pv)
        if p=="": raise ConversionError(f"第 {pos} 行地块字段为空。")
        out.append(PointRow(pos,p,opt_float(get(r,fields["area"])),clean(get(r,fields["land"])),
                            std_forest(get(r,fields["forest"])),std_grade(get(r,fields["grade"])),
                            as_float(px,"POINT_X",pos),as_float(py,"POINT_Y",pos)))
    if not out: raise ConversionError("没有可处理的数据行。")
    return out,resolved

def detect_zone(rows,expected=None):
    zones=set()
    for r in rows:
        x=abs(r.point_x)
        if x<10_000_000:
            raise ConversionError(f"POINT_X={r.point_x:.3f} 看起来是不带带号的 Easting，请先在 ArcGIS 使用 Project 转为 Zone N 后重新添加 XY。")
        z=int(x//1_000_000)
        if not 20<=z<=60: raise ConversionError(f"POINT_X={r.point_x:.3f} 无法识别合理带号。")
        if not 0<=abs(r.point_y)<=10_000_000: raise ConversionError(f"POINT_Y={r.point_y:.3f} 量级异常。")
        zones.add(z)
    if len(zones)!=1: raise ConversionError(f"初表中检测到多个带号：{sorted(zones)}")
    z=next(iter(zones))
    if expected is not None and z!=expected: raise ConversionError(f"检测到 Zone {z}，与 --expected-zone {expected} 不一致。")
    return z

def same_point(a,b,tol):
    return math.hypot(a.point_x-b.point_x,a.point_y-b.point_y)<=tol

def split_rings(points,tol):
    rings=[]; cur=[]
    for p in points:
        cur.append(p)
        if len(cur)>=4 and same_point(cur[0],p,tol):
            unique={(round(x.point_x,6),round(x.point_y,6)) for x in cur[:-1]}
            if len(unique)<3: raise ConversionError(f"地块 {p.parcel} 出现少于 3 个唯一顶点的面环。")
            rings.append(cur); cur=[]
    if cur:
        raise ConversionError(f"地块 {points[0].parcel} 存在未闭合面环，最后记录行号 {cur[-1].source_row}。")
    return rings

def build_output_rows(rows,tol):
    grouped=OrderedDict()
    for r in rows: grouped.setdefault(r.parcel,[]).append(r)
    out=[]; summary=OrderedDict()
    for parcel,pts in grouped.items():
        rings=split_rings(pts,tol)
        for ring_no,ring in enumerate(rings,start=1):
            for i,p in enumerate(ring):
                corner="J1" if i==len(ring)-1 else f"J{i+1}"
                out.append(OutputRow(p.source_row,parcel,ring_no,corner,p.area,p.land,p.forest,p.grade,p.point_x,p.point_y))
        first=pts[0]
        summary[parcel]={"area":first.area,"land":first.land,"forest":first.forest,"grade":first.grade,
                         "records":len(pts),"rings":len(rings)}
    return out,summary

def default_output_path(inp):
    stem=inp.stem.replace("_初表","_拐点坐标表")
    if stem==inp.stem: stem += "_拐点坐标表"
    return inp.with_name(stem+".xlsx")

def usage_type(path):
    n=path.stem
    if "永久" in n: return "永久"
    if "临时" in n: return "临时"
    return ""

def write_workbook(path,rows,summary,zone,small_threshold,b8_rows):
    try: import xlsxwriter
    except ImportError as e: raise ConversionError("生成 Excel 需要 XlsxWriter，请安装 requirements.txt。") from e
    path.parent.mkdir(parents=True,exist_ok=True)
    wb=xlsxwriter.Workbook(str(path))
    header=wb.add_format({"bold":True,"font_color":"#FFFFFF","bg_color":"#2F75B5","align":"center","valign":"vcenter","border":1})
    center=wb.add_format({"align":"center","valign":"vcenter","border":1})
    coord=wb.add_format({"align":"center","valign":"vcenter","border":1,"num_format":"0.000"})
    area4=wb.add_format({"align":"center","valign":"vcenter","border":1,"num_format":"0.0000"})
    title=wb.add_format({"bold":True,"font_size":16,"align":"center","valign":"vcenter"})
    label=wb.add_format({"bold":True,"bg_color":"#EAF2F8","border":1})
    meta=wb.add_format({"border":1})
    note=wb.add_format({"italic":True,"font_color":"#5F6B76","text_wrap":True})
    mono=wb.add_format({"font_name":"Consolas","font_size":10})
    start_blue=wb.add_format({"bold":True,"bg_color":"#EDF4FB","align":"center","border":1,"top":2,"top_color":"#5B9BD5"})
    start_yellow=wb.add_format({"bold":True,"bg_color":"#FFF4CC","align":"center","border":1,"top":2,"top_color":"#D6B656"})
    start_blue_coord=wb.add_format({"bold":True,"bg_color":"#EDF4FB","align":"center","border":1,"top":2,"top_color":"#5B9BD5","num_format":"0.000"})
    start_yellow_coord=wb.add_format({"bold":True,"bg_color":"#FFF4CC","align":"center","border":1,"top":2,"top_color":"#D6B656","num_format":"0.000"})

    ws=wb.add_worksheet("拐点坐标表")
    ws.merge_range("A1:D1","使用林地拐点坐标表",title)
    meta_rows=[("坐标系","2000国家大地坐标系（CGCS2000）","分带","3°"),
               ("投影类型","高斯-克吕格","带号",str(zone)),
               ("计算单位","米","坐标精度","0.001 m"),
               ("坐标说明","纵坐标 X = 北坐标（POINT_Y）","横坐标 Y","东坐标（POINT_X）")]
    for rr,row in enumerate(meta_rows,start=2):
        for cc,v in enumerate(row): ws.write(rr,cc,v,label if cc in (0,2) else meta)
    for c,h in enumerate(["地块号","拐点号","纵坐标（X）","横坐标（Y）"]): ws.write(7,c,h,header)
    lp=lr=None
    for idx,r in enumerate(rows,start=8):
        if r.parcel!=lp: f1,f2=start_blue,start_blue_coord
        elif r.ring!=lr: f1,f2=start_yellow,start_yellow_coord
        else: f1,f2=center,coord
        ws.write(idx,0,r.parcel,f1); ws.write(idx,1,r.corner,f1)
        ws.write_number(idx,2,r.northing_x,f2); ws.write_number(idx,3,r.easting_y,f2)
        lp,lr=r.parcel,r.ring
    multi=[str(p) for p,s in summary.items() if s["rings"]>1]
    text=f"说明：共{len(summary)}个地块、{sum(s['rings'] for s in summary.values())}个独立闭合面环。各面环按原始顶点顺序编号，末点重复 J1 表示闭合；同一地块存在多个面环时，每个面环均从 J1 重新编号。"
    if multi: text+=" 多面环地块："+"、".join(multi)+"。"
    nr=9+len(rows); ws.merge_range(nr,0,nr,3,text,note)
    ws.set_column("A:A",10); ws.set_column("B:B",12); ws.set_column("C:D",20); ws.freeze_panes(8,0)

    txt=wb.add_worksheet("报批文本")
    txt.merge_range("A1:B1","拐点坐标文本（可直接复制到报批坐标文件）",wb.add_format({"bold":True,"font_size":14,"align":"center"}))
    attrs=[("[属性描述]",""),("坐标系","2000国家大地坐标系"),("几度分带","3"),("投影类型","高斯克吕格"),
           ("计算单位","米"),("带号",str(zone)),("精度","0.001")]
    for rr,(a,b) in enumerate(attrs,start=2): txt.write(rr,0,a,label); txt.write(rr,1,b)
    txt.write(10,0,"[地块坐标]",label); rr=11; prev=None
    for r in rows:
        key=(r.parcel,r.ring)
        if prev is not None and key!=prev: rr+=1
        txt.write(rr,0,f"{r.corner},{r.parcel},{r.northing_x:.3f},{r.easting_y:.3f}",mono); rr+=1; prev=key
    txt.set_column("A:A",46); txt.set_column("B:B",30); txt.freeze_panes(11,0)

    ck=wb.add_worksheet("核查辅助")
    headers=["原始行号","地块号","面环号","拐点号","面积(ha)","地类","森林类别","保护等级","POINT_X(Easting)","POINT_Y(Northing)","说明"]
    for c,h in enumerate(headers): ck.write(0,c,h,header)
    last={(r.parcel,r.ring):r.source_row for r in rows}
    for rr,r in enumerate(rows,start=1):
        notes=[]
        if r.source_row==last[(r.parcel,r.ring)]: notes.append("J1闭合点")
        if r.area is not None and r.area<=small_threshold: notes.append("面积极小，请复核")
        vals=[r.source_row,r.parcel,r.ring,r.corner,r.area,r.land,r.forest,r.grade,r.point_x,r.point_y,"；".join(notes)]
        for c,v in enumerate(vals): ck.write(rr,c,v,area4 if c==4 else center)
    for c,w in enumerate([10,10,10,10,12,14,14,12,20,20,22]): ck.set_column(c,c,w)
    ck.freeze_panes(1,0)

    sm=wb.add_worksheet("地块汇总")
    headers=["地块号","面积(ha)","地类","森林类别","保护等级","拐点记录数","面环数","核查"]
    for c,h in enumerate(headers): sm.write(0,c,h,header)
    for rr,(p,s) in enumerate(summary.items(),start=1):
        check="面积极小，请复核" if s["area"] is not None and s["area"]<=small_threshold else "正常"
        vals=[p,s["area"],s["land"],s["forest"],s["grade"],s["records"],s["rings"],check]
        for c,v in enumerate(vals): sm.write(rr,c,v,area4 if c==1 else center)
    for c,w in enumerate([10,12,14,14,12,14,10,18]): sm.set_column(c,c,w)
    sm.freeze_panes(1,0)

    b8=wb.add_worksheet("附表B.8三栏版")
    bt=wb.add_format({"bold":True,"font_size":14,"font_name":"宋体","align":"center","valign":"vcenter"})
    bh=wb.add_format({"font_size":10,"font_name":"宋体","align":"center","valign":"vcenter","text_wrap":True,"border":1})
    bc=wb.add_format({"font_size":9.5,"font_name":"宋体","align":"center","valign":"vcenter","border":1})
    bn=wb.add_format({"font_size":9.5,"font_name":"宋体","align":"center","valign":"vcenter","border":1,"num_format":"0.000"})
    for c in (0,3,6): b8.set_column(c,c,9)
    for c in (1,2,4,5,7,8): b8.set_column(c,c,18)
    per=max(10,int(b8_rows)); per_section=per*3; sections=math.ceil(len(rows)/per_section); breaks=[]
    for sec in range(sections):
        base=sec*(per+3)
        b8.merge_range(base,0,base,8,"表B.8  项目使用林地小班位置拐点坐标表",bt)
        for c,h in enumerate(["地块\n序号","x","y"]*3): b8.write(base+1,c,h,bh)
        for ri in range(per):
            sr=base+2+ri
            for block in range(3):
                idx=sec*per_section+block*per+ri; c0=block*3
                if idx<len(rows):
                    r=rows[idx]; b8.write(sr,c0,r.parcel,bc); b8.write_number(sr,c0+1,r.northing_x,bn); b8.write_number(sr,c0+2,r.easting_y,bn)
                else:
                    for c in range(c0,c0+3): b8.write_blank(sr,c,None,bc)
        if sec<sections-1: breaks.append(base+per+3)
    if breaks: b8.set_h_pagebreaks(breaks)
    b8.set_portrait(); b8.set_paper(9); b8.fit_to_pages(1,0); b8.set_margins(0.25,0.25,0.35,0.35)
    wb.close()

def report_data(inp,out,rows,summary,zone,resolved,threshold):
    return {"input":str(inp),"output":str(out),"zone":zone,"parcels":len(summary),
            "rings":sum(s["rings"] for s in summary.values()),"point_records":len(rows),
            "multi_ring_parcels":[p for p,s in summary.items() if s["rings"]>1],
            "small_area_parcels":[p for p,s in summary.items() if s["area"] is not None and s["area"]<=threshold],
            "resolved_fields":resolved}

def parse_args(argv=None):
    p=argparse.ArgumentParser(description="ArcGIS 林勘拐点初表 → 报批 Excel")
    p.add_argument("input"); p.add_argument("--output"); p.add_argument("--expected-zone",type=int)
    p.add_argument("--parcel-field"); p.add_argument("--small-area-threshold",type=float,default=0.0001)
    p.add_argument("--closure-tolerance",type=float,default=0.001); p.add_argument("--b8-rows",type=int,default=50)
    p.add_argument("--check-only",action="store_true"); p.add_argument("--json",action="store_true")
    return p.parse_args(argv)

def main(argv=None):
    a=parse_args(argv); inp=Path(a.input).expanduser().resolve()
    if not inp.exists(): print(f"错误：文件不存在：{inp}",file=sys.stderr); return 2
    out=Path(a.output).expanduser().resolve() if a.output else default_output_path(inp)
    try:
        headers,raw=read_table(inp); parsed,resolved=parse_rows(headers,raw,a.parcel_field)
        zone=detect_zone(parsed,a.expected_zone); rows,summary=build_output_rows(parsed,a.closure_tolerance)
        rep=report_data(inp,out,rows,summary,zone,resolved,a.small_area_threshold)
        if not a.check_only: write_workbook(out,rows,summary,zone,a.small_area_threshold,a.b8_rows)
        if a.json: print(json.dumps(rep,ensure_ascii=False,indent=2))
        else:
            print(f"[OK] 输入：{inp}"); print(f"[OK] 带号：Zone {zone}"); print(f"[OK] 地块数：{rep['parcels']}")
            print(f"[OK] 独立闭合面环：{rep['rings']}"); print(f"[OK] 拐点记录：{rep['point_records']}")
            if rep["multi_ring_parcels"]: print("[INFO] 多面环地块："+"、".join(map(str,rep["multi_ring_parcels"])))
            if rep["small_area_parcels"]: print("[WARNING] 面积极小，请复核："+"、".join(map(str,rep["small_area_parcels"])))
            if not a.check_only: print(f"[OK] 输出：{out}")
        return 0
    except ConversionError as e:
        print(f"错误：{e}",file=sys.stderr); return 2

if __name__=="__main__":
    raise SystemExit(main())
